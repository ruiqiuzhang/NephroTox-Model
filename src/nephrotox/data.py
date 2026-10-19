"""Read the sole, frozen DATA source without rewriting any chemical structure."""
from __future__ import annotations

from collections import Counter
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook
from rdkit import Chem

ROOT = Path(__file__).resolve().parents[2]


def rows_fingerprint(rows):
    encoded = json.dumps(rows, ensure_ascii=False, sort_keys=True,
                         separators=(',', ':'), allow_nan=False).encode('utf-8')
    return hashlib.sha256(encoded).hexdigest()


def load_data(path=None, manifest_path=None):
    """Return (rows, source_info), preserving all 12 nonempty source columns.

    The workbook SHA locks this release to the supplied 1673-record DATA file.
    Five completely empty unnamed Excel columns are ignored after verifying
    every cell in them is empty. No historical source or chemistry service is
    accessed. Specified and unspecified stereo are both preserved verbatim.
    """
    root = ROOT
    path = Path(path) if path else root / 'data/NephroTox_Modeling_Data.xlsx'
    manifest_path = Path(manifest_path) if manifest_path else root / 'data/source_manifest.json'
    expected = json.loads(manifest_path.read_text(encoding='utf-8'))
    file_hash = hashlib.sha256(path.read_bytes()).hexdigest()
    if file_hash != expected['sha256']:
        raise ValueError('DATA SHA256 differs from this release: no alternate or historical source is accepted')
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        if workbook.sheetnames != [expected['sheet_name']]:
            raise ValueError('Unexpected DATA worksheets')
        iterator = workbook[expected['sheet_name']].iter_rows(values_only=True)
        raw_headers = list(next(iterator))
        active = [i for i, name in enumerate(raw_headers) if name is not None]
        headers = [raw_headers[i] for i in active]
        if headers != expected['source_columns'] or len(set(headers)) != len(headers):
            raise ValueError('DATA column names/order differ from the frozen source schema')
        unnamed = set(range(len(raw_headers))) - set(active)
        if len(unnamed) != expected['empty_unnamed_columns_ignored']:
            raise ValueError('Unexpected unnamed source columns')
        rows = []
        for excel_row, values in enumerate(iterator, 2):
            if any(values[i] is not None for i in unnamed):
                raise ValueError(f'Nonempty unnamed source column at Excel row {excel_row}')
            if not any(v is not None for v in values):
                raise ValueError(f'Unexpected empty record at Excel row {excel_row}')
            row = {raw_headers[i]: values[i] for i in active}
            if any(v is None or v == '' for v in row.values()):
                raise ValueError(f'Missing DATA field at Excel row {excel_row}')
            if type(row['label']) is not int or row['label'] not in (0, 1):
                raise ValueError(f'Nonbinary label at Excel row {excel_row}')
            if type(row['PubChem_CID']) is not int or row['PubChem_CID'] <= 0:
                raise ValueError(f'Invalid PubChem CID at Excel row {excel_row}')
            if not isinstance(row['compound_id'], str) or not isinstance(row['standardized_smiles'], str):
                raise ValueError(f'Invalid identity/SMILES type at Excel row {excel_row}')
            molecule = Chem.MolFromSmiles(row['standardized_smiles'])
            if molecule is None or len(Chem.GetMolFrags(molecule)) != 1:
                raise ValueError(f'Invalid or multicomponent molecule: {row["compound_id"]}')
            if any(atom.GetAtomicNum() == 0 for atom in molecule.GetAtoms()):
                raise ValueError(f'Unspecified atom in molecule: {row["compound_id"]}')
            rows.append(row)
    finally:
        workbook.close()
    if len(rows) != expected['n_rows'] or len({r['compound_id'] for r in rows}) != len(rows):
        raise ValueError('DATA population or unique compound IDs differ from the release')
    counts = {str(k): v for k, v in Counter(r['label'] for r in rows).items()}
    if counts != expected['label_counts']:
        raise ValueError('DATA positive/negative counts differ from the release')
    info = dict(expected)
    info['rows_fingerprint'] = rows_fingerprint(rows)
    info['source_file'] = 'data/NephroTox_Modeling_Data.xlsx'
    info['all_source_values_preserved'] = True
    return rows, info
